import io
from dataclasses import dataclass
from pathlib import PurePath

import fitz
from bs4 import BeautifulSoup
from docx import Document as DocxDocument


class ParseError(ValueError):
    """Raised when a supported document cannot be converted to policy text."""


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    text: str
    format: str


_SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".html", ".htm", ".md", ".markdown", ".txt"}


def _decode_text(payload: bytes, extension: str) -> str:
    try:
        return payload.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ParseError(f"{extension} content must be valid UTF-8 text.") from error


def _parse_html(payload: bytes) -> str:
    soup = BeautifulSoup(_decode_text(payload, ".html"), "html.parser")
    for node in soup(["script", "style", "noscript"]):
        node.decompose()
    return soup.get_text("\n", strip=True)


def _parse_docx(payload: bytes) -> str:
    try:
        document = DocxDocument(io.BytesIO(payload))
    except Exception as error:
        raise ParseError("DOCX content could not be parsed.") from error
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)


def _parse_pdf(payload: bytes) -> str:
    try:
        document = fitz.open(stream=payload, filetype="pdf")
    except Exception as error:
        raise ParseError("PDF content could not be parsed.") from error
    try:
        return "\n".join(page.get_text("text").strip() for page in document)
    finally:
        document.close()


def parse_document(filename: str, payload: bytes) -> ParsedDocument:
    """Parse only supported policy data; parsed text is evidence, never executable instructions."""

    extension = PurePath(filename).suffix.lower()
    if extension not in _SUPPORTED_EXTENSIONS:
        raise ParseError(f"Unsupported document extension: {extension or '<none>'}.")
    if not payload:
        raise ParseError("Document payload is empty.")

    if extension in {".txt", ".md", ".markdown"}:
        text = _decode_text(payload, extension)
    elif extension in {".html", ".htm"}:
        text = _parse_html(payload)
    elif extension == ".docx":
        text = _parse_docx(payload)
    else:
        text = _parse_pdf(payload)

    if not text.strip():
        raise ParseError("Document contains no extractable text.")
    return ParsedDocument(text=text, format=extension.removeprefix("."))
